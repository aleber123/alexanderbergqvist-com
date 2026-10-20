"""Generate the free timesheet templates served from /downloads/.

    ~/Downloads/asc-launcher/.venv/bin/python scripts/gen-tidrapport-mall.py

Writes:
  public/downloads/tidrapport-mall-vecka.xlsx  – weekly, with formulas
  public/downloads/tidrapport-mall-manad.xlsx  – monthly, OB kväll/natt/helg,
                                                 övertid, frånvaro, pay summary
  public/downloads/tidrapport-mall-manad.csv   – monthly layout, plain CSV
  public/downloads/tidrapport-mall-manad.pdf   – monthly, printable A4

Needs openpyxl + reportlab. Re-run after changing the layout; the files
are committed so the site build doesn't depend on Python.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'public', 'downloads')
os.makedirs(OUT, exist_ok=True)

INK = '1F1B16'
BLUE = '1976D2'
LIGHT = 'EAF2FB'
thin = Side(style='thin', color='C9C2B8')
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def xlsx_week():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tidrapport vecka'
    ws.sheet_view.showGridLines = False

    widths = {'A': 12, 'B': 13, 'C': 10, 'D': 10, 'E': 11, 'F': 15, 'G': 13, 'H': 34}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws['A1'] = 'Tidrapport – vecka'
    ws['A1'].font = Font(size=18, bold=True, color=INK)
    ws['A2'] = 'Fyll i de ljusblå fälten. Timmar, summor och lön räknas ut automatiskt.'
    ws['A2'].font = Font(italic=True, color='6B645B')

    fields = [
        ('A4', 'Namn', 'C4'),
        ('A5', 'Arbetsgivare', 'C5'),
        ('A6', 'Veckans måndag', 'C6'),
        ('E4', 'Timlön (kr)', 'G4'),
        ('E5', 'OB-tillägg (%)', 'G5'),
        ('E6', 'Vecka', 'G6'),
    ]
    inp = PatternFill('solid', fgColor=LIGHT)
    for label_cell, label, value_cell in fields:
        ws[label_cell] = label
        ws[label_cell].font = Font(bold=True)
        c = ws[value_cell]
        c.fill = inp
        c.border = box
    ws.merge_cells('C4:D4')
    ws.merge_cells('C5:D5')
    ws.merge_cells('C6:D6')
    ws['C6'].number_format = 'yyyy-mm-dd'
    ws['G4'].number_format = '#,##0.00'
    ws['G5'].number_format = '0'
    ws['G5'] = 50
    ws['G6'] = '=IF(C6="","",ISOWEEKNUM(C6))'
    ws['G6'].fill = PatternFill(None)

    head = ['Dag', 'Datum', 'Start', 'Slut', 'Rast (min)', 'Arbetade timmar', 'varav OB-tim', 'Anteckning']
    r0 = 8
    for i, h in enumerate(head):
        c = ws.cell(row=r0, column=i + 1, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = box
    ws.row_dimensions[r0].height = 30

    days = ['Måndag', 'Tisdag', 'Onsdag', 'Torsdag', 'Fredag', 'Lördag', 'Söndag']
    for i, d in enumerate(days):
        r = r0 + 1 + i
        ws.cell(row=r, column=1, value=d).font = Font(bold=True)
        ws.cell(row=r, column=2, value=f'=IF($C$6="","",$C$6+{i})').number_format = 'yyyy-mm-dd'
        for col in (3, 4):
            c = ws.cell(row=r, column=col)
            c.number_format = 'hh:mm'
            c.fill = inp
        ws.cell(row=r, column=5).fill = inp
        # MOD(slut-start;1) handles night shifts that end after midnight.
        ws.cell(row=r, column=6, value=f'=IF(OR(C{r}="",D{r}=""),"",ROUND(MOD(D{r}-C{r},1)*24-N(E{r})/60,2))').number_format = '0.00'
        ws.cell(row=r, column=7).fill = inp
        ws.cell(row=r, column=7).number_format = '0.00'
        ws.cell(row=r, column=8).fill = inp
        for col in range(1, 9):
            ws.cell(row=r, column=col).border = box
            if col < 8:
                ws.cell(row=r, column=col).alignment = Alignment(horizontal='center')

    rs = r0 + 8
    first, last = r0 + 1, r0 + 7
    ws.cell(row=rs, column=1, value='Summa').font = Font(bold=True)
    ws.cell(row=rs, column=6, value=f'=SUM(F{first}:F{last})').number_format = '0.00'
    ws.cell(row=rs, column=7, value=f'=SUM(G{first}:G{last})').number_format = '0.00'
    for col in range(1, 9):
        c = ws.cell(row=rs, column=col)
        c.font = Font(bold=True)
        c.border = box
        c.fill = PatternFill('solid', fgColor='F3EFE8')
        if col < 8:
            c.alignment = Alignment(horizontal='center')

    r = rs + 2
    rows = [
        ('Grundlön (timmar × timlön)', f'=IF($G$4="","",F{rs}*$G$4)'),
        ('OB-tillägg (OB-tim × timlön × OB-%)', f'=IF($G$4="","",G{rs}*$G$4*$G$5/100)'),
        ('Summa före skatt', f'=IF($G$4="","",F{r}+F{r + 1})'),
    ]
    for i, (label, formula) in enumerate(rows):
        ws.cell(row=r + i, column=1, value=label).font = Font(bold=(i == 2))
        ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=5)
        c = ws.cell(row=r + i, column=6, value=formula)
        c.number_format = '#,##0.00 "kr"'
        c.font = Font(bold=(i == 2))
        c.border = box

    r += 5
    ws.cell(row=r, column=1, value='Underskrift anställd: ______________________      Attesterat av: ______________________')
    r += 2
    ws.cell(row=r, column=1, value='Tips: skriv tider som 07:30 och 16:00. Slutar passet efter midnatt räknas det ändå rätt.').font = Font(italic=True, color='6B645B')
    ws.cell(row=r + 1, column=1, value='Slipp mallen: appen Tidrapportera räknar timmar, OB och lön automatiskt – alexanderbergqvist.com/tidrapport/').font = Font(italic=True, color=BLUE)

    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='600', allow_blank=True,
                        error='Ange rast i minuter, t.ex. 30.', errorTitle='Rast')
    ws.add_data_validation(dv)
    dv.add(f'E{first}:E{last}')

    ws.print_area = f'A1:H{r + 1}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = 'A9'

    # No cached values are stored, so ask the app to recalc on open.
    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = 'Tidrapport mall – vecka'
    wb.properties.creator = 'Alexander Bergqvist'
    path = os.path.join(OUT, 'tidrapport-mall-vecka.xlsx')
    wb.save(path)
    return path


def xlsx_month():
    """Monthly template: one row per day, OB split into kväll/natt/helg,
    overtime and absence, with a pay summary that takes each OB level as
    either % of the hourly wage or a fixed kr/tim amount."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tidrapport månad'
    ws.sheet_view.showGridLines = False

    widths = {'A': 12, 'B': 7, 'C': 9, 'D': 9, 'E': 9, 'F': 11, 'G': 10, 'H': 10, 'I': 10,
              'J': 10, 'K': 14, 'L': 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    ws['A1'] = 'Tidrapport – månad'
    ws['A1'].font = Font(size=18, bold=True, color=INK)
    ws['A2'] = ('Fyll i de ljusblå fälten. Datum, veckodag, timmar och lön räknas ut automatiskt. '
                'OB-timmarna skriver du själv enligt ditt avtal.')
    ws['A2'].font = Font(italic=True, color='6B645B')

    inp = PatternFill('solid', fgColor=LIGHT)

    def field(label_cell, label, value_cell, fmt=None, value=None):
        ws[label_cell] = label
        ws[label_cell].font = Font(bold=True)
        c = ws[value_cell]
        c.fill = inp
        c.border = box
        if fmt:
            c.number_format = fmt
        if value is not None:
            c.value = value

    field('A4', 'Namn', 'C4')
    field('A5', 'Arbetsgivare', 'C5')
    field('A6', 'Månad (valfritt datum i månaden)', 'C6', 'yyyy-mm')
    for rng in ('C4:E4', 'C5:E5', 'C6:E6'):
        ws.merge_cells(rng)
    ws['A6'].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[6].height = 30

    # Rates. OB and overtime: % of hourly wage, OR kr/tim (kr wins if filled).
    ws['G3'] = 'Lön och tillägg'
    ws['G3'].font = Font(bold=True, color=BLUE)
    ws['I3'] = '%'
    ws['J3'] = 'eller kr/tim'
    for c in ('I3', 'J3'):
        ws[c].font = Font(bold=True)
        ws[c].alignment = Alignment(horizontal='center')
    field('G4', 'Timlön (kr)', 'I4', '#,##0.00')
    ws.merge_cells('I4:J4')
    rate_rows = [(5, 'OB kväll', 50), (6, 'OB natt', 70), (7, 'OB helg', 100), (8, 'Övertid', 50)]
    for r, label, pct in rate_rows:
        field(f'G{r}', label, f'I{r}', '0', pct)
        cj = ws[f'J{r}']
        cj.fill = inp
        cj.border = box
        cj.number_format = '#,##0.00'
    ws['G9'] = 'Procentsatserna är exempel – byt till ditt avtals.'
    ws['G9'].font = Font(italic=True, size=9, color='6B645B')

    head = ['Datum', 'Dag', 'Start', 'Slut', 'Rast (min)', 'Arbetade tim', 'OB kväll tim',
            'OB natt tim', 'OB helg tim', 'Övertid tim', 'Frånvaro', 'Anteckning']
    r0 = 11
    for i, h in enumerate(head):
        c = ws.cell(row=r0, column=i + 1, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = box
    ws.row_dimensions[r0].height = 30

    first = r0 + 1
    first_of_month = 'DATE(YEAR($C$6),MONTH($C$6),1)'
    for i in range(31):
        r = first + i
        # Rows past the month's last day stay blank (29–31).
        ws.cell(row=r, column=1,
                value=f'=IF($C$6="","",IF(MONTH({first_of_month}+{i})=MONTH($C$6),{first_of_month}+{i},""))'
                ).number_format = 'yyyy-mm-dd'
        # Locale-independent weekday (TEXT(...;"ddd") depends on Excel's language).
        ws.cell(row=r, column=2,
                value=f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"mån","tis","ons","tor","fre","lör","sön"))')
        for col in (3, 4):
            c = ws.cell(row=r, column=col)
            c.number_format = 'hh:mm'
            c.fill = inp
        ws.cell(row=r, column=5).fill = inp
        # MOD(slut-start;1) handles night shifts that end after midnight.
        ws.cell(row=r, column=6,
                value=f'=IF(OR(C{r}="",D{r}=""),"",ROUND(MOD(D{r}-C{r},1)*24-N(E{r})/60,2))'
                ).number_format = '0.00'
        for col in range(7, 13):
            ws.cell(row=r, column=col).fill = inp
        for col in range(7, 11):
            ws.cell(row=r, column=col).number_format = '0.00'
        for col in range(1, 13):
            ws.cell(row=r, column=col).border = box
            if col < 12:
                ws.cell(row=r, column=col).alignment = Alignment(horizontal='center')
    last = first + 30

    rs = last + 1
    ws.cell(row=rs, column=1, value='Summa').font = Font(bold=True)
    for col, letter in zip(range(6, 11), 'FGHIJ'):
        ws.cell(row=rs, column=col, value=f'=SUM({letter}{first}:{letter}{last})').number_format = '0.00'
    for col in range(1, 13):
        c = ws.cell(row=rs, column=col)
        c.font = Font(bold=True)
        c.border = box
        c.fill = PatternFill('solid', fgColor='F3EFE8')
        if col < 12:
            c.alignment = Alignment(horizontal='center')

    def add(hours_cell, rate_row):
        return (f'=IF($I$4="","",IF($J${rate_row}<>"",{hours_cell}*$J${rate_row},'
                f'{hours_cell}*$I$4*N($I${rate_row})/100))')

    r = rs + 2
    rows = [
        ('Grundlön (arbetade timmar × timlön)', f'=IF($I$4="","",F{rs}*$I$4)'),
        ('OB-tillägg kväll', add(f'G{rs}', 5)),
        ('OB-tillägg natt', add(f'H{rs}', 6)),
        ('OB-tillägg helg', add(f'I{rs}', 7)),
        ('Övertidstillägg', add(f'J{rs}', 8)),
        ('Summa före skatt', f'=IF($I$4="","",SUM(F{r}:F{r + 4}))'),
    ]
    for i, (label, formula) in enumerate(rows):
        bold = i == len(rows) - 1
        ws.cell(row=r + i, column=1, value=label).font = Font(bold=bold)
        ws.merge_cells(start_row=r + i, start_column=1, end_row=r + i, end_column=5)
        c = ws.cell(row=r + i, column=6, value=formula)
        c.number_format = '#,##0.00 "kr"'
        c.font = Font(bold=bold)
        c.border = box
    note_r = r + len(rows)
    ws.cell(row=note_r, column=1,
            value='Semesterersättning och skatt ingår inte. Övertidstillägget räknas utöver grundlönen för timmarna.'
            ).font = Font(italic=True, size=9, color='6B645B')

    r = note_r + 2
    ws.cell(row=r, column=1, value='Underskrift anställd: ______________________      Attesterat av: ______________________')
    r += 2
    ws.cell(row=r, column=1, value='Tips: skriv tider som 07:30 och 16:00. Slutar passet efter midnatt räknas det ändå rätt.').font = Font(italic=True, color='6B645B')
    ws.cell(row=r + 1, column=1, value='Slipp mallen: appen Tidrapportera räknar OB och lön åt dig för varje pass – alexanderbergqvist.com/tidrapport/').font = Font(italic=True, color=BLUE)

    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='600', allow_blank=True,
                        error='Ange rast i minuter, t.ex. 30.', errorTitle='Rast')
    ws.add_data_validation(dv)
    dv.add(f'E{first}:E{last}')
    dv2 = DataValidation(type='list', formula1='"Sjuk,VAB,Semester,Ledig,Föräldraledig,Annat"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add(f'K{first}:K{last}')

    ws.print_area = f'A1:L{r + 1}'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.freeze_panes = f'A{first}'

    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = 'Tidrapport mall – månad'
    wb.properties.creator = 'Alexander Bergqvist'
    path = os.path.join(OUT, 'tidrapport-mall-manad.xlsx')
    wb.save(path)
    return path


def csv_month():
    """Plain CSV of the monthly layout for apps that can't open .xlsx.
    Semicolon + UTF-8 BOM so Swedish Excel splits columns and shows å/ä/ö."""
    path = os.path.join(OUT, 'tidrapport-mall-manad.csv')
    head = ['Dag', 'Datum', 'Start', 'Slut', 'Rast (min)', 'Arbetade tim', 'OB kväll tim',
            'OB natt tim', 'OB helg tim', 'Övertid tim', 'Frånvaro', 'Anteckning']
    lines = [';'.join(head)]
    for d in range(1, 32):
        lines.append(';'.join([str(d)] + [''] * (len(head) - 1)))
    lines.append(';'.join(['Summa'] + [''] * (len(head) - 1)))
    with open(path, 'w', encoding='utf-8-sig', newline='') as f:
        f.write('\r\n'.join(lines) + '\r\n')
    return path


def pdf_month():
    path = os.path.join(OUT, 'tidrapport-mall-manad.pdf')
    c = canvas.Canvas(path, pagesize=A4)
    c.setTitle('Tidrapport mall – månad')
    c.setAuthor('Alexander Bergqvist')
    W, H = A4
    left, right = 12 * mm, W - 12 * mm

    c.setFont('Helvetica-Bold', 18)
    c.drawString(left, H - 18 * mm, 'Tidrapport – månad')

    # Header fields
    c.setFont('Helvetica', 9)
    y = H - 28 * mm
    fields = [('Namn', left, 84 * mm), ('Arbetsgivare', left + 96 * mm, 90 * mm)]
    fields2 = [('Månad och år', left, 40 * mm), ('Timlön (kr)', left + 46 * mm, 38 * mm),
               ('Anställningsnr / avdelning', left + 96 * mm, 90 * mm)]
    for row_y, row in ((y, fields), (y - 11 * mm, fields2)):
        for label, x, w in row:
            c.setFillColor(colors.HexColor('#6B645B'))
            c.drawString(x, row_y + 5 * mm, label)
            c.setStrokeColor(colors.HexColor('#9C948A'))
            c.line(x, row_y, x + w, row_y)
    c.setFillColor(colors.black)

    cols = [('Dag', 9), ('Start', 14), ('Slut', 14), ('Rast', 11), ('Arb. tim', 15),
            ('OB kväll', 15), ('OB natt', 15), ('OB helg', 15), ('Övertid', 15),
            ('Frånvaro', 17), ('Anteckning', 0)]
    total_w = right - left
    fixed = sum(w for _, w in cols) * mm
    widths = [w * mm if w else total_w - fixed for _, w in cols]
    xs = [left]
    for w in widths:
        xs.append(xs[-1] + w)

    top = y - 19 * mm
    row_h = 6.0 * mm
    # Header row
    c.setFillColor(colors.HexColor('#1976D2'))
    c.rect(left, top - row_h, total_w, row_h, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont('Helvetica-Bold', 7.5)
    for (label, _), x, w in zip(cols, xs, widths):
        c.drawCentredString(x + w / 2, top - row_h + 2 * mm, label)

    c.setFont('Helvetica', 8)
    c.setStrokeColor(colors.HexColor('#C9C2B8'))
    for i in range(31):
        yy = top - row_h * (i + 2)
        if i % 2 == 1:
            c.setFillColor(colors.HexColor('#F6F3EE'))
            c.rect(left, yy, total_w, row_h, stroke=0, fill=1)
        c.setFillColor(colors.black)
        c.drawCentredString(xs[0] + widths[0] / 2, yy + 2 * mm, str(i + 1))
    # Summa row
    ys = top - row_h * 33
    c.setFillColor(colors.HexColor('#EAE4DA'))
    c.rect(left, ys, total_w, row_h, stroke=0, fill=1)
    c.setFillColor(colors.black)
    c.setFont('Helvetica-Bold', 7.5)
    c.drawString(left + 1 * mm, ys + 2 * mm, 'Summa')

    # Grid
    bottom = ys
    for x in xs:
        c.line(x, top, x, bottom)
    for i in range(34):
        c.line(left, top - row_h * i, right, top - row_h * i)

    # OB / overtime rates used (so the reader can check the sums)
    ry = bottom - 9 * mm
    c.setFont('Helvetica', 8.5)
    c.setFillColor(colors.HexColor('#6B645B'))
    c.drawString(left, ry + 4 * mm, 'Tillägg enligt avtal (% av timlön eller kr/tim):')
    c.setFillColor(colors.black)
    for j, label in enumerate(('OB kväll', 'OB natt', 'OB helg', 'Övertid')):
        x = left + j * 46 * mm
        c.drawString(x, ry - 2 * mm, label)
        c.setStrokeColor(colors.HexColor('#9C948A'))
        c.line(x + 15 * mm, ry - 2.5 * mm, x + 40 * mm, ry - 2.5 * mm)

    # Signatures
    sy = ry - 20 * mm
    c.setFont('Helvetica', 9)
    for label, x in (('Underskrift anställd', left), ('Attesterat av arbetsgivare', left + 96 * mm)):
        c.line(x, sy, x + 84 * mm, sy)
        c.setFillColor(colors.HexColor('#6B645B'))
        c.drawString(x, sy - 4.5 * mm, label + '  (datum och namnförtydligande)')
        c.setFillColor(colors.black)

    c.setFont('Helvetica-Oblique', 7.5)
    c.setFillColor(colors.HexColor('#6B645B'))
    c.drawString(left, 12 * mm,
                 'Arb. tim = slut - start - rast. Rast i minuter. Fler mallar: alexanderbergqvist.com/tidrapport/tidrapport-mall/')
    c.drawString(left, 8 * mm,
                 'Slipp mallen: appen Tidrapportera räknar OB och lön åt dig för varje pass och exporterar samma rapport som PDF.')
    c.showPage()
    c.save()
    return path


if __name__ == '__main__':
    for p in (xlsx_week(), xlsx_month(), csv_month(), pdf_month()):
        print(p, os.path.getsize(p), 'bytes')
